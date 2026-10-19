"""Dashboard Đối chiếu DO vs Hóa đơn — Streamlit + Odoo XML-RPC real-time."""
from __future__ import annotations

import io
import xmlrpc.client
import socket
from collections import defaultdict
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

socket.setdefaulttimeout(30)

URL = st.secrets["odoo"]["url"]
DB = st.secrets["odoo"]["db"]
USER = st.secrets["odoo"]["user"]
PWD = st.secrets["odoo"]["password"]
ODOO_CONTEXT = {"allowed_company_ids": [1], "lang": "vi_VN", "tz": "Asia/Bangkok"}
LIMIT = 200
COSTING_LABELS = {"average": "AVCO", "fifo": "FIFO", "standard": "Standard"}


def search_read_all(mdl, uid: int, model: str, domain: list, fields: list) -> list:
    offset = 0
    results = []
    while True:
        batch = mdl.execute_kw(
            DB, uid, PWD, model, "search_read", [domain],
            {"fields": fields, "limit": LIMIT, "offset": offset, "context": ODOO_CONTEXT},
        )
        results.extend(batch)
        if len(batch) < LIMIT:
            break
        offset += LIMIT
    return results


@st.cache_data(ttl=300, show_spinner="Đang tải dữ liệu từ Odoo...")
def fetch_data() -> list[dict]:
    common = xmlrpc.client.ServerProxy(f"{URL}/xmlrpc/2/common", allow_none=True)
    uid = common.authenticate(DB, USER, PWD, {})
    if not uid:
        st.error("Xác thực Odoo thất bại. Kiểm tra cấu hình Secrets.")
        return []
    mdl = xmlrpc.client.ServerProxy(f"{URL}/xmlrpc/2/object", allow_none=True)

    sales = search_read_all(mdl, uid, "sale.order", [
        ["state", "in", ["sale", "done"]],
    ], ["id", "name", "partner_id", "order_line"])

    if not sales:
        return []

    sol_ids = []
    for so in sales:
        sol_ids.extend(so["order_line"])

    so_lines = search_read_all(mdl, uid, "sale.order.line", [
        ["id", "in", sol_ids],
        ["product_id", "!=", False],
    ], ["id", "order_id", "product_id", "product_uom_qty", "name"])
    sol_map = {sol["id"]: sol for sol in so_lines}

    stock_moves = search_read_all(mdl, uid, "stock.move", [
        ["sale_line_id", "in", sol_ids],
        ["state", "=", "done"],
    ], ["id", "sale_line_id", "product_id", "quantity", "picking_id", "value", "price_unit"])

    do_qty_by_sol: dict[int, float] = defaultdict(float)
    cost_by_sol: dict[int, float] = defaultdict(float)
    for sm in stock_moves:
        sol_id = sm["sale_line_id"][0] if sm["sale_line_id"] else None
        if sol_id:
            do_qty_by_sol[sol_id] += sm["quantity"]
            cost_by_sol[sol_id] += abs(sm.get("value", 0) or 0)

    inv_lines = search_read_all(mdl, uid, "account.move.line", [
        ["sale_line_ids", "in", sol_ids],
        ["parent_state", "=", "posted"],
        ["display_type", "not in", ["line_section", "line_note"]],
    ], ["id", "sale_line_ids", "quantity", "move_id", "product_id"])

    move_ids = list({il["move_id"][0] for il in inv_lines if il["move_id"]})
    move_types = {}
    if move_ids:
        moves = search_read_all(mdl, uid, "account.move", [
            ["id", "in", move_ids],
        ], ["id", "move_type"])
        move_types = {m["id"]: m["move_type"] for m in moves}

    inv_qty_by_sol: dict[int, float] = defaultdict(float)
    for il in inv_lines:
        move_id = il["move_id"][0] if il["move_id"] else None
        move_type = move_types.get(move_id, "")
        sign = -1.0 if move_type == "out_refund" else 1.0
        for sol_id in il.get("sale_line_ids", []):
            if sol_id in sol_map:
                inv_qty_by_sol[sol_id] += il["quantity"] * sign

    product_ids = list({sol["product_id"][0] for sol in so_lines if sol["product_id"]})
    products = search_read_all(mdl, uid, "product.product", [
        ["id", "in", product_ids],
    ], ["id", "name", "categ_id", "standard_price", "uom_id"])
    product_map = {p["id"]: p for p in products}

    categ_ids = list({p["categ_id"][0] for p in products if p["categ_id"]})
    categs = search_read_all(mdl, uid, "product.category", [
        ["id", "in", categ_ids],
    ], ["id", "name", "property_cost_method"])
    categ_map = {c["id"]: c for c in categs}

    so_map = {so["id"]: so for so in sales}

    rows = []
    for sol in so_lines:
        sol_id = sol["id"]
        so = so_map.get(sol["order_id"][0], {}) if sol["order_id"] else {}
        product = product_map.get(sol["product_id"][0], {}) if sol["product_id"] else {}
        categ = categ_map.get(product.get("categ_id", [False])[0], {}) if product.get("categ_id") else {}

        qty_do = do_qty_by_sol.get(sol_id, 0.0)
        qty_inv = inv_qty_by_sol.get(sol_id, 0.0)
        if qty_do == 0 and qty_inv == 0:
            continue

        qty_diff = qty_inv - qty_do
        if qty_do > 0 and cost_by_sol.get(sol_id, 0) > 0:
            unit_cost = cost_by_sol[sol_id] / qty_do
        else:
            unit_cost = product.get("standard_price", 0.0)

        prod_name = product.get("name", "")
        if isinstance(prod_name, dict):
            prod_name = prod_name.get("vi_VN", prod_name.get("en_US", str(prod_name)))

        rows.append({
            "Mã SO": so.get("name", ""),
            "Khách hàng": (so.get("partner_id", [False, ""])[1] if so.get("partner_id") else ""),
            "Sản phẩm": prod_name,
            "ĐVT": (product.get("uom_id", [False, ""])[1] if product.get("uom_id") else ""),
            "PP Giá vốn": COSTING_LABELS.get(categ.get("property_cost_method", ""), "N/A"),
            "SL DO": qty_do,
            "SL HĐ": qty_inv,
            "Chênh lệch KL": qty_diff,
            "ĐG vốn BQ": unit_cost,
            "GV theo DO": qty_do * unit_cost,
            "GV theo HĐ": qty_inv * unit_cost,
            "Chênh lệch GV": qty_diff * unit_cost,
        })

    return rows


def build_excel(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Đối chiếu DO vs HĐ"

    header_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = "BÁO CÁO ĐỐI CHIẾU PHIẾU GIAO (DO) vs HÓA ĐƠN"
    title_cell.font = Font(name="Arial", bold=True, size=14, color="2F5496")
    title_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:L2")
    ws["A2"].value = f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="666666")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        ("Mã SO", 12), ("Khách hàng", 30), ("Sản phẩm", 30), ("ĐVT", 8),
        ("PP Giá vốn", 12), ("SL Phiếu giao (DO)", 18), ("SL Hóa đơn", 15),
        ("Chênh lệch KL", 16), ("ĐG vốn BQ", 16), ("GV theo DO", 18),
        ("GV theo HĐ", 18), ("Chênh lệch GV", 18),
    ]
    for col_idx, (name, width) in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    col_keys = ["Mã SO", "Khách hàng", "Sản phẩm", "ĐVT", "PP Giá vốn",
                "SL DO", "SL HĐ", "Chênh lệch KL", "ĐG vốn BQ",
                "GV theo DO", "GV theo HĐ", "Chênh lệch GV"]
    data_font = Font(name="Arial", size=10)
    red_font = Font(name="Arial", size=10, color="CC0000")
    green_font = Font(name="Arial", size=10, color="008000")

    for row_idx, (_, row) in enumerate(df.iterrows(), 5):
        has_diff = abs(row["Chênh lệch KL"]) > 0.001
        for col_idx, key in enumerate(col_keys, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[key])
            cell.border = thin_border
            cell.font = data_font
            if col_idx in (6, 7, 8):
                cell.number_format = "#,##0.000"
            elif col_idx in (9, 10, 11, 12):
                cell.number_format = "#,##0"
            if col_idx == 8 and has_diff:
                cell.font = red_font if row["Chênh lệch KL"] < 0 else green_font
            if col_idx == 12 and has_diff:
                cell.font = red_font if row["Chênh lệch GV"] < 0 else green_font

    total_row = 5 + len(df)
    total_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    ws.merge_cells(f"A{total_row}:K{total_row}")
    label = ws.cell(row=total_row, column=1, value="TỔNG CHÊNH LỆCH GIÁ VỐN")
    label.font = Font(name="Arial", bold=True, size=11)
    label.fill = total_fill
    label.alignment = Alignment(horizontal="right")
    label.border = thin_border
    for c in range(2, 12):
        ws.cell(row=total_row, column=c).fill = total_fill
        ws.cell(row=total_row, column=c).border = thin_border

    total_val = df["Chênh lệch GV"].sum()
    tc = ws.cell(row=total_row, column=12, value=total_val)
    tc.font = Font(name="Arial", bold=True, size=12, color="CC0000" if total_val < 0 else "008000")
    tc.fill = total_fill
    tc.number_format = "#,##0"
    tc.border = thin_border

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:L{total_row - 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    st.set_page_config(page_title="Đối chiếu DO vs HĐ", page_icon="📊", layout="wide")

    rows = fetch_data()

    if not rows:
        st.warning("Không có dữ liệu.")
        return

    df = pd.DataFrame(rows)
    excel_bytes = build_excel(df)
    fname = f"DoiChieu_DO_vs_HoaDon_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    import base64
    b64 = base64.b64encode(excel_bytes).decode()
    auto_download = f'''
        <html><body>
        <script>
            var link = document.createElement("a");
            link.href = "data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{b64}";
            link.download = "{fname}";
            document.body.appendChild(link);
            link.click();
        </script>
        <p style="font-family:Arial; font-size:16px; text-align:center; margin-top:40px;">
            ✅ Đang tải file <b>{fname}</b>...<br><br>
            Nếu file chưa tải, <a href="data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{b64}" download="{fname}">bấm vào đây</a>.
        </p>
        </body></html>
    '''
    st.components.v1.html(auto_download, height=150)


if __name__ == "__main__":
    main()
